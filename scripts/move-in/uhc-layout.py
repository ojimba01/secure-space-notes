"""Reads UHC's Move-In Supports Checklist (.xlsx) and writes the item list the
app fills it from: src/lib/uhcMoveIn/layout.ts.

Every line the app shows carries the cells it is written back to, so the
spreadsheet UHC receives is their own file with the answers in place. When UHC
reissues the form, put the new file at public/form-templates/
uhc-move-in-supports.xlsx, run this again, and look over the diff.

    pip install openpyxl
    python3 scripts/move-in/uhc-layout.py

The sheet has no formulas: its totals are typed zeros. The app adds them up.
"""
import json, re, sys
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as C

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / 'public/form-templates/uhc-move-in-supports.xlsx'
OUT = ROOT / 'src/lib/uhcMoveIn/layout.ts'

wb = openpyxl.load_workbook(SRC)

def text(ws, ref):
    v = ws[ref].value
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''

def filled(ws, ref):
    f = ws[ref].fill
    return bool(f and f.fill_type)

def slug(s):
    return re.sub(r'[^a-z0-9]+', '_', s.lower()).strip('_')[:40]

# Labels that ask for a word or two beside the quantity (a flavour, a type).
def wants_note(label):
    l = label.lower()
    return bool(re.search(r'(:|-)\s*$', label)) or any(w in l for w in ('type', 'flavor', 'specify', 'list ', ' or ', 'hard/soft', 'front/back', '/beef/'))

def tidy(label):
    """The label as it reads on screen: no trailing colon or dash, no asterisk."""
    s = re.sub(r'^\*', '', label)
    s = re.sub(r'[\s:–-]+$', '', s)
    return s.strip()

def split_note(label):
    """'Yogurt Flavor' → ('Yogurt', 'Flavor'); 'Salad Dressing: Type' → ('Salad Dressing', 'Type')."""
    m = re.match(r'^(.*?)\s*(?::|\s[-–]\s?)\s*(.*)$', label)
    if label.lower() == 'flavor':
        return label, 'Which flavor'
    if m and m.group(1):
        note = m.group(2).strip() or 'Type'
        return m.group(1).strip(), note[0].upper() + note[1:]
    m = re.match(r'^(.*?)\s+(Flavor|Type)$', label, re.I)
    if m:
        return m.group(1).strip(), m.group(2).capitalize()
    if re.search(r'\bor\b', label, re.I):
        return label, 'Which one'
    return label, 'Type'

used_ids = set()
def uid(base):
    i, out = 1, base
    while out in used_ids:
        i += 1
        out = f'{base}_{i}'
    used_ids.add(out)
    return out

NO_NOTE = {'H20'}

# How a few cells read on screen where UHC's own words do not stand alone.
LABELS = {
    ('Food Items', 'B101'): 'Cereal',
    ('Household Items', 'H20'): 'Long dresser',
    ('Household Items', 'H21'): 'Tall dresser',
    ('Household Items', 'B21'): 'Standard height',
    ('Household Items', 'B22'): 'Low height',
}
HEADINGS = {
    ('Food Items', 'H33'): 'Cheese form',
    ('Food Items', 'E71'): 'Sausages',
    ('Food Items', 'E54'): 'Coffee/Filters',
}

def merged_span(ws, ref):
    for m in ws.merged_cells.ranges:
        if m.coord.split(':')[0] == ref:
            return m.min_col, m.max_col
    return None

def block_lines(ws, sheet_key, rows, triples, section_rows, stop_rows=(), group_override=None):
    """Lines in a grid of (qty, item, total) column triples.

    `section_rows` maps a row to a section title: the section starts there.
    A filled cell in an item column with items beneath it is a group heading
    (Cheddar: block / shredded …); its items carry it as their group.
    """
    sections = []
    cur = None
    for r in rows:
        if r in section_rows:
            cur = {'title': section_rows[r], 'lines': []}
            sections.append(cur)
            continue
        if r in stop_rows or cur is None:
            continue
        # A blank cell ends the group heading above it in that column.
        for _, ci, _ in triples:
            if not text(ws, f'{ci}{r}'):
                cur.get('_heads', {}).pop(ci, None)
        for col_q, col_i, col_t in triples:
            ref = f'{col_i}{r}'
            label = text(ws, ref)
            if not label:
                continue
            below = text(ws, f'{col_i}{r + 1}')
            is_heading = filled(ws, ref) and below and not filled(ws, f'{col_i}{r + 1}')
            if group_override and ref in group_override:
                is_heading = group_override[ref]
            if ws.title == 'Household Items' and ref == 'B18':
                cur.get('_heads', {}).pop('B', None)
            if is_heading:
                head = HEADINGS.get((ws.title, ref), tidy(label))
                span = merged_span(ws, ref)
                for _, ci, _ in triples:
                    if ci == col_i or (span and span[0] <= C(ci) <= span[1]):
                        cur.setdefault('_heads', {})[ci] = head
                continue
            group = cur.get('_heads', {}).get(col_i)
            shown = LABELS.get((ws.title, ref), tidy(label))
            line = {
                'id': uid(f'{sheet_key}.{slug((group + " " if group else "") + shown)}'),
                'label': shown,
                'cell': ref,
                'qty': f'{col_q}{r}' if col_q else None,
                'total': f'{col_t}{r}',
            }
            if group:
                line['group'] = group
            if (wants_note(label) or (group and label.lower() == 'flavor')) and ref not in NO_NOTE:
                line['note'] = True
                line['label'], line['noteLabel'] = split_note(shown)
                # The note is written after UHC's own words: "Yogurt Flavor: Strawberry".
                line['text'] = str(ws[ref].value).rstrip()
            cur['lines'].append(line)
    for s in sections:
        s.pop('_heads', None)
    return sections

def free_rows(prefix, rows):
    """Blank lines for anything not listed: typed item, quantity, total."""
    return [{'id': f'{prefix}{i + 1}', **r, 'free': True} for i, r in enumerate(rows)]

# ---- Food Items --------------------------------------------------------------
ws = wb['Food Items']
food_sections = {2: 'Vegetables'}
for r in range(3, 148):
    a = text(ws, f'A{r}')
    if a and filled(ws, f'A{r}'):
        food_sections[r] = a
food_triples = [('A', 'B', 'C'), ('D', 'E', 'F'), ('G', 'H', 'I')]
food = block_lines(ws, 'food', range(2, 148), food_triples, food_sections,
                   group_override={'H85': False, 'E53': False})
for s in food:
    if s['title'] == 'Other Food Items':
        s['lines'] = free_rows('food.other', [{'cell': 'B129', 'qty': 'A129', 'total': 'I129', 'multi': True}])
food = [s for s in food if s['lines']]

# ---- Pantry, hygiene, clothing ----------------------------------------------
ws = wb['Pantry_Hygiene']
pantry = block_lines(ws, 'pantry', range(1, 24), [('A', 'B', 'C'), ('D', 'E', 'F')],
                     {1: 'Kitchen', 6: 'Cleaning Supplies', 15: 'Office Supplies', 20: 'Extras'})
pantry.append({'title': 'Other Pantry Items', 'lines': free_rows('pantry.other', [
    {'cell': 'A25', 'qty': None, 'total': 'F25'}, {'cell': 'A26', 'qty': None, 'total': 'F26'}])})
hygiene = block_lines(ws, 'hygiene', range(1, 9), [('I', 'J', 'K'), ('L', 'M', 'N'), ('O', 'P', 'Q')],
                      {1: 'Personal Hygiene'})
hygiene.append({'title': 'Other Personal Hygiene', 'lines': free_rows('hygiene.other', [
    {'cell': 'J10', 'qty': 'I10', 'total': 'Q10'}, {'cell': 'J11', 'qty': 'I11', 'total': 'Q11'}])})
clothing_lines = []
for r in range(2, 7):
    label = text(ws, f'W{r}')
    clothing_lines.append({'id': uid(f'clothing.{slug(label)}'), 'label': tidy(label), 'cell': f'W{r}',
                           'qty': f'U{r}', 'total': f'V{r}', 'size': f'T{r}',
                           **({'note': True, 'noteLabel': 'Front/Back', 'text': str(ws[f'W{r}'].value).rstrip()} if wants_note(label) else {})})
    if clothing_lines[-1]['label'] and wants_note(label):
        clothing_lines[-1]['label'] = re.sub(r'\s*\(.*\)$', '', clothing_lines[-1]['label'])
clothing = [{'title': 'Clothing', 'lines': clothing_lines},
            {'title': 'Other Clothing', 'lines': free_rows('clothing.other', [
                {'cell': f'W{r}', 'qty': f'U{r}', 'total': f'V{r}', 'size': f'T{r}'} for r in range(8, 12)])}]

# ---- Household Items ---------------------------------------------------------
ws = wb['Household Items']
household = block_lines(ws, 'household', range(1, 26), [('A', 'B', 'C'), ('D', 'E', 'F'), ('G', 'H', 'I')],
                        {1: 'Bathroom Items', 5: 'Kitchen Items', 13: 'Bedroom Items', 23: 'Furniture'},
                        group_override={'H20': False, 'B20': True, 'E20': True})
# "Kitchen Furniture" is a heading in H23 beside "Living Room Furniture".
household.append({'title': 'Other Household Items', 'lines': free_rows('household.other', [
    {'cell': f'B{r}', 'qty': f'A{r}', 'total': f'C{r}'} for r in range(27, 31)])})
caps = []
for r in range(2, 32):
    m = re.match(r'^(.*?)\.{2,}\s*([\d,]+\.\d\d)', text(ws, f'L{r}'))
    if m:
        caps.append({'label': m.group(1).strip(), 'max': float(m.group(2).replace(',', ''))})

# ---- Services ----------------------------------------------------------------
ws = wb['Services']
def svc(ref, total, note=None, label=None):
    lab = label or tidy(text(ws, ref))
    out = {'id': uid(f'services.{slug(lab)}'), 'label': lab, 'cell': ref, 'qty': None, 'total': total}
    if note:
        out['note'] = True
        out['noteLabel'] = note
        out['text'] = str(ws[ref].value).rstrip()
    return out
services = [
    {'title': 'Residence Fees', 'lines': [
        svc('B3', 'C3', 'Property name', label='Application Fee'),
        # Rows 4–6 are blank on UHC's sheet, though its note asks for a W-9 with
        # a security deposit. The deposit is written into the first of them.
        svc('B4', 'C4', label='Security Deposit'),
        svc('B7', 'C7', 'What the deposit is for', label='Other Home Deposit'),
    ]},
    {'title': 'Utilities Deposits', 'lines': [svc(f'B{r}', f'C{r}', 'Which utility' if r == 16 else None) for r in range(9, 17)]},
    {'title': 'Moving Company', 'lines': [svc('B17', 'C17', 'Company name', label='Moving Company')]},
    {'title': 'Identification Documents', 'lines': [svc(f'B{r}', f'C{r}', 'Which document' if r == 23 else None) for r in range(20, 24)]},
]
# Lines whose amounts need the landlord's W-9, invoices and a lease.
needs_paperwork = [l['id'] for l in services[0]['lines'] + services[1]['lines'] + services[2]['lines']]

# ---- Member Information --------------------------------------------------------
member = {
    'provider': 'C6', 'caseManager': 'C7', 'cmPhone': 'C8', 'cmEmail': 'C9',
    'memberName': 'C11', 'medicaidId': 'C12', 'householdSize': 'C13', 'newAddress': 'C14',
    'memberPhone': 'C15', 'moveInDate': 'C16', 'emergencyContact': 'B18', 'delivery': 'B20',
}

tabs = [
    {'id': 'services', 'title': 'Services', 'sheet': 'xl/worksheets/sheet2.xml', 'totalCell': 'C28', 'sections': services},
    {'id': 'food', 'title': 'Food', 'sheet': 'xl/worksheets/sheet3.xml', 'totalCell': 'I148', 'sections': food},
    {'id': 'pantry', 'title': 'Pantry', 'sheet': 'xl/worksheets/sheet4.xml', 'totalCell': 'F27', 'sections': pantry},
    {'id': 'hygiene', 'title': 'Hygiene', 'sheet': 'xl/worksheets/sheet4.xml', 'totalCell': 'Q12', 'sections': hygiene},
    {'id': 'clothing', 'title': 'Clothing', 'sheet': 'xl/worksheets/sheet4.xml', 'totalCell': 'W17', 'sections': clothing},
    {'id': 'household', 'title': 'Household', 'sheet': 'xl/worksheets/sheet5.xml', 'totalCell': 'G31', 'sections': household},
]

def strip(o):
    if isinstance(o, dict):
        return {k: strip(v) for k, v in o.items() if v is not None}
    if isinstance(o, list):
        return [strip(v) for v in o]
    return o

count = sum(len(s['lines']) for t in tabs for s in t['sections'])
ts = f"""// Generated by scripts/move-in/uhc-layout.py from UHC's Move-In Supports
// Checklist. Do not edit by hand: change the script and run it again.
//
// Each line names the cells it is written to in UHC's own spreadsheet
// (public/form-templates/uhc-move-in-supports.xlsx). {count} lines.
import type {{ UhcTab }} from './types';

export const MEMBER_SHEET = 'xl/worksheets/sheet1.xml';
export const MEMBER_CELLS = {json.dumps(member, indent=2)} as const;
export const GRAND_TOTAL_CELL = 'F22';

export const UHC_TABS: UhcTab[] = {json.dumps(strip(tabs), indent=2)};

/** UHC's furniture price caps. Amounts over a cap are reviewed one by one, not refused. */
export const FURNITURE_CAPS: {{ label: string; max: number }}[] = {json.dumps(caps, indent=2)};

/** Lines whose amount needs the landlord's W-9, invoices and a lease (or a letter). */
export const NEEDS_PAPERWORK: string[] = {json.dumps(needs_paperwork, indent=2)};
"""
OUT.write_text(ts)
print(f'{count} lines → {OUT.relative_to(ROOT)}')
